"""Extract the monthly cases of Swissdec appendix 1 into a TypeScript data file.

Run: uv run --with openpyxl python3 scripts/extract-swissdec-appendix.py
Input:  sources/rules/Anhang_1_QST_Berechnung_Monat_20260306.xlsx
Output: src/domain/source-tax/monthly/swissdec-appendix.generated.ts, and a skip report on stdout.

Each case sheet lists one person for a tax year. The script reads the rows of every month column:
code, rate, social-insurance days, the wages, the rate-defining salary and the tax amount.
A sheet is skipped when it is out of scope or when a cell does not map to the model of `MonthInput`.
"""

import json
import re
from fractions import Fraction
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "sources/rules/Anhang_1_QST_Berechnung_Monat_20260306.xlsx"
OUTPUT = ROOT / "src/domain/source-tax/monthly/swissdec-appendix.generated.ts"

MONTHS = ["Jan", "Feb", "Mrz", "Apr", "Mai", "Jun", "Jul", "Aug", "Sep", "Okt", "Nov", "Dez"]

# Sheets that are out of scope by title (spec.md § 12: no correction, no canton change, no annual model).
SCOPE_SKIPS = {
    "M14.1": "corrections of March and August (Korrektur) and back payments after leaving in April and September",
    "M19.1": "foreign working days are split off (Ausscheidung Auslandtage); the model has no foreign-days input",
    "M31": "foreign working days are split off (Ausscheidung Auslandtage)",
    "M32": "payment in the following year and foreign working days",
    "M33": "left before the tax year, payments in the following year, foreign working days",
    "M33.1": "canton change and foreign working days",
    "M35": "correction of the leaving month (Korrektur)",
    "M36": "left before the tax year, payment in the following year",
    "M37": "payment after leaving with a new tariff code in the payment month",
    "M38": "left before the tax year, payment in the following year",
    "M38.1": "foreign working days and payment in the following year",
    "M40": "correction of earlier months (Korrektur)",
    "M41": "correction of earlier months (Korrektur)",
    "M42": "correction of earlier months (Korrektur)",
    "YM43": "canton and system change from the annual model to the monthly model",
    "MY44": "canton and system change from the monthly model to the annual model",
    "M46": "canton change",
}


def number(value):
    return value if isinstance(value, (int, float)) and not isinstance(value, bool) else None


def to_rappen(value, what):
    """CHF float to integer Rappen. Raise when the cell is not a whole number of Rappen."""
    scaled = value * 100
    if abs(scaled - round(scaled)) > 1e-6:
        raise Unsupported(f"{what} {value} is not a whole number of Rappen (the sheet does not round a pro-rata 13th month salary)")
    return round(scaled)


class Unsupported(Exception):
    pass


def extract(ws):
    rows = {}
    for row in ws.iter_rows():
        label = row[0].value
        if isinstance(label, str):
            rows.setdefault(label.strip(), []).append([c.value for c in row[1:]])

    header = rows.get("Bezeichnung")
    if header is None:
        raise Unsupported("no month header row")
    columns = [i for i, v in enumerate(header[0]) if v in MONTHS]
    if not columns:
        raise Unsupported("the sheet lists weeks, not month columns")

    def first(prefix, exact=False):
        for label, lists in rows.items():
            if (label == prefix) if exact else label.startswith(prefix):
                return lists[0]
        return None

    def many(prefix):
        out = []
        for label, lists in rows.items():
            if label.startswith(prefix) and "Total" not in label and "kumul" not in label:
                out.append(lists[0])
        return out

    def cell(values, i):
        return values[i] if values is not None and i < len(values) else None

    def pick(lists, i):
        for values in lists:
            v = number(cell(values, i))
            if v is not None:
                return v
        return None

    code_row = first("QST-Tarifcode")
    sv_row = first("SV-Tage", exact=True)
    sb_row = first("QST-SB-Lohn", exact=True)
    sb_aperiodic_row = first("QST-SB-aperiodisch", exact=True)
    if code_row is None or sv_row is None or sb_row is None:
        raise Unsupported("a required row is missing (code, SV days or rate-defining salary)")
    rate_rows = many("QST-Satz")
    tax_rows = many("QST-Betrag")
    lohn_rows = many("QST-Lohn")
    total_lohn = first("QST-Lohn Total", exact=True)
    thirteenth_row = next((v[0] for k, v in rows.items() if re.match(r"^13\. Monatslohn", k)), None)
    weitere = first("Weitere Beschäftigung (X)", exact=True)
    bg_own = first("BG AG 1", exact=True)
    bg_other = first("BG andere AG", exact=True)
    bg_total = first("BG Total", exact=True)
    hours_row = first("Anzahl Stunden", exact=True)
    qst_hours = first("QST-Stunden", exact=True)
    pay_row = first("Stundenlohn", exact=True)
    entry_row = first("Eintritt", exact=True)
    exit_row = first("Austritt", exact=True)
    if not rate_rows or not tax_rows or not (lohn_rows or total_lohn):
        raise Unsupported("rate, tax or QST wage row is missing")

    # Entry and exit dates sit in the column of their month. A date in the column of its own
    # month belongs to the tax year. Other dates (for example an entry in an earlier year) do not count.
    monthly_hours_row = first("Monats-Arbeitzeit AG 1", exact=True)

    def dated(row):
        out = []
        for j, v in enumerate(row or []):
            if hasattr(v, "year") and j < len(header[0]) and header[0][j] in MONTHS:
                if MONTHS.index(header[0][j]) + 1 == v.month:
                    out.append(v)
        return out

    years = {d.year for d in dated(entry_row) + dated(exit_row)}
    if not years:
        raise Unsupported("no entry or exit date in a month column, so the tax year is unknown")
    if len(years) > 1:
        raise Unsupported(f"the dates of the sheet name several tax years: {sorted(years)}")
    (year,) = years

    def spells(row):
        return {(v.year, v.month): v for v in (row or []) if hasattr(v, "year")}

    entries = spells(entry_row)
    exits = spells(exit_row)

    def percent(value, what):
        """A level from the sheet as an exact fraction of 1, limited to 4 decimals."""
        fraction = Fraction(str(value)).limit_denominator(10000)
        if abs(float(fraction) - value) > 1e-9:
            raise Unsupported(f"{what} {value} is not an exact decimal")
        return fraction

    months, skipped_months = [], []
    for i in columns:
        month = MONTHS.index(header[0][i]) + 1
        code = cell(code_row, i)
        sv_days = number(cell(sv_row, i))
        if code is None or sv_days is None:
            continue
        try:
            months.append(extract_month(i, month, year, code, sv_days, locals()))
        except Unsupported as error:
            skipped_months.append({"month": month, "reason": str(error)})

    if not months:
        raise Unsupported("no month can be tested: " + "; ".join(f"{m['month']}: {m['reason']}" for m in skipped_months))
    return {
        "sheet": ws.title,
        "title": ws["A1"].value.strip(),
        "year": year,
        "months": months,
        "skippedMonths": skipped_months,
    }


def extract_month(i, month, year, code, sv_days, scope):
    cell, pick = scope["cell"], scope["pick"]
    code = str(code).replace("AON", "A0N")
    rate = pick(scope["rate_rows"], i)
    tax = pick(scope["tax_rows"], i)
    lohn = pick(scope["lohn_rows"], i)
    total_lohn = scope["total_lohn"]
    if lohn is None and total_lohn is not None:
        lohn = number(cell(total_lohn, i))
    sb = number(cell(scope["sb_row"], i))
    if rate is None or lohn is None or sb is None:
        raise Unsupported("a rate, wage or rate-defining cell is empty")
    scaled_rate = rate * 10000
    if abs(scaled_rate - round(scaled_rate)) > 1e-6:
        raise Unsupported(f"rate {rate} has more than two decimals as a percentage")
    aperiodic = number(cell(scope["sb_aperiodic_row"], i)) or 0
    thirteenth = number(cell(scope["thirteenth_row"], i)) or 0
    lohn_r = to_rappen(lohn, "QST-Lohn")
    aperiodic_r = to_rappen(aperiodic, "aperiodic")
    thirteenth_r = to_rappen(thirteenth, "13th month")

    entry = scope["entries"].get((year, month))
    leave = scope["exits"].get((year, month))
    if entry is not None and leave is not None:
        if leave < entry:
            raise Unsupported("two spells of employment in one month")
        period = {"kind": "startAndEnd", "start": [year, month, entry.day], "end": [year, month, leave.day]}
    elif entry is not None:
        period = {"kind": "start", "date": [year, month, entry.day]}
    elif leave is not None:
        period = {"kind": "end", "date": [year, month, leave.day]}
    else:
        period = {"kind": "fullMonth"}

    # Levels are exact percent values as fractions {num, den}. A level from hours is hours / monthly hours.
    percent = scope["percent"]
    bg_own, bg_other, bg_total = scope["bg_own"], scope["bg_other"], scope["bg_total"]
    hours_row, monthly_hours_row = scope["hours_row"], scope["monthly_hours_row"]
    other = number(cell(bg_other, i)) or 0
    own_fraction = None
    if monthly_hours_row is not None and number(cell(hours_row, i)):
        own_fraction = Fraction(int(number(cell(hours_row, i))), int(number(cell(monthly_hours_row, i))))
        sheet_own = number(cell(bg_own, i))
        if sheet_own is None or abs(float(own_fraction) - sheet_own) > 6e-5:
            raise Unsupported("the level hours / monthly hours differs from the sheet level")
    elif number(cell(bg_own, i)) is not None:
        own_fraction = percent(number(cell(bg_own, i)), "own level")

    def as_percent(fraction):
        fraction = fraction * 100
        return {"num": fraction.numerator, "den": fraction.denominator}

    employment = {"kind": "single"}
    if other > 0:
        if own_fraction is None:
            raise Unsupported("the own employment level is missing")
        total_fraction = own_fraction + percent(other, "level of the other job")
        if total_fraction > 1:
            raise Unsupported(f"the total employment level {float(total_fraction) * 100:.2f}% is above 100% (the model allows at most 100%)")
        employment = {"kind": "levels", "ownPercent": as_percent(own_fraction), "totalPercent": as_percent(total_fraction)}
    elif scope["weitere"] is not None and cell(scope["weitere"], i) not in (None, ""):
        # Another job exists, its level is unknown: the own level is converted to 100% (KS 45 § 6.4).
        if own_fraction is None:
            raise Unsupported("the own employment level is missing")
        employment = {"kind": "ownLevelOnly", "ownPercent": as_percent(own_fraction)}

    result = {
        "month": month,
        "code": code,
        "rate": round(scaled_rate),
        "svDays": round(sv_days),
        "period": period,
        "employment": employment,
        "aperiodic": aperiodic_r,
        "thirteenthMonth": thirteenth_r,
        "expectedRateDefining": f"{sb:.4f}".rstrip("0").rstrip(".") if sb != round(sb, 2) else f"{sb:.2f}",
        "expectedTax": None if tax is None else round(tax * 100, 4),
    }
    if scope["qst_hours"] is not None:
        # Hourly pay, weekly payment: the monthly sum is converted to 180 hours (KS 45 § 6.5).
        hours = number(cell(hours_row, i))
        if hours is None or hours == 0:
            raise Unsupported("no hours")
        result["hourly"] = {"pay": lohn_r - aperiodic_r, "hours": round(hours)}
        result["periodic"] = 0
    else:
        result["periodic"] = lohn_r - aperiodic_r - thirteenth_r
    return result


def case_line(case):
    """One case: the header on one line, then one line for each month."""
    head = {k: v for k, v in case.items() if k not in ("months", "skippedMonths")}
    months = ",\n".join("\t\t" + json.dumps(m, ensure_ascii=False) for m in case["months"])
    return "\t{ " + ", ".join(f"{k}: {json.dumps(v, ensure_ascii=False)}" for k, v in head.items()) + f", months: [\n{months},\n\t] }},"


def main():
    wb = openpyxl.load_workbook(SOURCE, data_only=True)
    cases, skipped = [], {}
    for ws in wb.worksheets:
        if not re.match(r"^(Y?M|MY)\d", ws.title):
            continue
        if ws.title in SCOPE_SKIPS:
            skipped[ws.title] = SCOPE_SKIPS[ws.title]
            continue
        try:
            cases.append(extract(ws))
        except Unsupported as error:
            skipped[ws.title] = str(error)

    lines = [
        "// Generated by scripts/extract-swissdec-appendix.py from",
        "// sources/rules/Anhang_1_QST_Berechnung_Monat_20260306.xlsx. Do not edit.",
        "// Amounts are Rappen. Rates are hundredths of a percent. `expectedRateDefining` is CHF as the sheet shows it.",
        "",
        "export type AppendixPeriod =",
        '\t| { kind: "fullMonth" }',
        '\t| { kind: "start"; date: [number, number, number] }',
        '\t| { kind: "end"; date: [number, number, number] }',
        '\t| { kind: "startAndEnd"; start: [number, number, number]; end: [number, number, number] };',
        "",
        "/** A percent value as an exact fraction. */",
        "export type AppendixPercent = { num: number; den: number };",
        "",
        "export type AppendixMonth = {",
        "\tmonth: number;",
        "\tcode: string;",
        "\trate: number;",
        "\tsvDays: number;",
        "\tperiod: AppendixPeriod;",
        "\temployment:",
        '\t\t| { kind: "single" }',
        '\t\t| { kind: "levels"; ownPercent: AppendixPercent; totalPercent: AppendixPercent }',
        '\t\t| { kind: "ownLevelOnly"; ownPercent: AppendixPercent };',
        "\tperiodic: number;",
        "\tthirteenthMonth: number;",
        "\taperiodic: number;",
        "\thourly?: { pay: number; hours: number };",
        "\texpectedRateDefining: string;",
        "\t/** Rappen, as the sheet computes it (no rounding). */",
        "\texpectedTax: number | null;",
        "};",
        "",
        "export type AppendixCase = { sheet: string; title: string; year: number; months: AppendixMonth[] };",
        "",
        "export const SWISSDEC_APPENDIX_MONTHLY: AppendixCase[] = [",
        *[case_line(c) for c in cases],
        "];",
        "",
    ]
    OUTPUT.write_text("\n".join(lines), encoding="utf-8")
    report = {
        "included": [c["sheet"] for c in cases],
        "skippedSheets": skipped,
        "skippedMonths": {c["sheet"]: c["skippedMonths"] for c in cases if c["skippedMonths"]},
    }
    print(json.dumps(report, ensure_ascii=False, indent=1))


if __name__ == "__main__":
    sys.exit(main())
